from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Отчёт_анализ_и_рефакторинг_RPM_PR_1.docx"
SCREENSHOTS = [
    (
        ROOT / "qa" / "photo_1_2026-09-15_16-54-10.jpg",
        "Исходный интерфейс приложения с данными студентов",
        "Рисунок 1 – Исходный интерфейс приложения до рефакторинга",
    ),
    (
        ROOT / "qa" / "photo_2_2026-09-15_16-54-10.jpg",
        "Исправленный интерфейс с данными, индикатором состояния и пагинацией",
        "Рисунок 2 – Исправленная версия с подключённой MySQL и пагинацией",
    ),
    (
        ROOT / "qa" / "photo_3_2026-09-15_16-54-10.jpg",
        "Исправленный интерфейс при пустом наборе данных",
        "Рисунок 3 – Отображение пустой таблицы в исправленной версии",
    ),
]


def set_run_font(run, name="Times New Roman", size=14, bold=False, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        borders.append(element)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    keep_with_next(paragraph)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=15 if level == 1 else 14, bold=True)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(4)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        text = text[len(bold_lead):]
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(1.25)
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(text), size=13)
    return paragraph


def add_code(doc, lines, caption):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    properties.append(shading)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, name="Consolas", size=9)
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(6)
    set_run_font(label.add_run(caption), size=11)


def add_figure(doc, image_path, alt_text, caption_text, width=15.2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(image_path), width=Cm(width))
    picture._inline.docPr.set("descr", alt_text)
    picture._inline.docPr.set("title", caption_text)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_after = Pt(8)
    set_run_font(caption.add_run(caption_text), size=11)


def add_table(doc, headers, rows, widths=None, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(str(text)), size=font_size, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, "EAF2F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(paragraph.add_run(str(value)), size=font_size, bold=str(value).startswith("Итого"))
            if widths:
                cell.width = Cm(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(3)
section.right_margin = Cm(1.5)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(14)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

title = doc.add_paragraph()
title.style = doc.styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Cm(4)
title.paragraph_format.space_after = Cm(1.5)
set_run_font(title.add_run("ОТЧЁТ ПО ПРОЕКТНОМУ ЗАНЯТИЮ"), size=18, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(subtitle.add_run("Анализ рефакторинг и оптимизация чужого проекта"), size=16, bold=True)

project = doc.add_paragraph()
project.alignment = WD_ALIGN_PARAGRAPH.CENTER
project.paragraph_format.space_before = Cm(1)
set_run_font(project.add_run("Проект RPM PR 1\nПриложение учёта студентов и успеваемости"), size=14)

branch = doc.add_paragraph()
branch.alignment = WD_ALIGN_PARAGRAPH.CENTER
branch.paragraph_format.space_before = Cm(1)
set_run_font(branch.add_run("Исходная ветка ошибки человечества\nИсправленная ветка codex/fixed-human-errors"), size=13)

author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.RIGHT
author.paragraph_format.left_indent = Cm(8.5)
author.paragraph_format.space_before = Cm(2)
author.paragraph_format.space_after = Pt(0)
set_run_font(
    author.add_run(
        "Выполнил: [ФИО студента]\n"
        "Группа: [номер группы]\n"
        "Проверил: [ФИО преподавателя]"
    ),
    size=13,
)

year = doc.add_paragraph()
year.alignment = WD_ALIGN_PARAGRAPH.CENTER
year.paragraph_format.space_before = Cm(3.5)
set_run_font(year.add_run(str(date.today().year)), size=14)
doc.add_page_break()

add_body(doc, "Тема: анализ качества кода, рефакторинг и оптимизация JavaFX-приложения, полученного от другой команды.", "Тема:")
add_body(doc, "Цель: найти фактические ошибки в чужом проекте, классифицировать их, выполнить исправления без изменения предметной логики и подготовить доказательства для защиты.", "Цель:")

add_heading(doc, "1 Ход работы")
add_heading(doc, "1.1 Исходные данные", 2)
add_body(doc, "Для анализа использован клонированный репозиторий RPM_PR_1. Исходной точкой является ветка «ошибки-человечества», коммит f253bf8. Чтобы не изменять переданный вариант, исправления выполнены в отдельной ветке codex/fixed-human-errors.")
add_body(doc, "Приложение ведёт учёт студентов, учебных предметов и оценок. Интерфейс реализован на JavaFX, доступ к данным — через JDBC и MySQL, схема базы создаётся миграциями Flyway.")

add_heading(doc, "1.2 Требования методички", 2)
add_body(doc, "Этапы 3–5 требуют найти и зафиксировать ошибки в чужом коде, выполнить приёмы Extract Method, Extract Class, Introduce Constant, Rename, инкапсуляцию и устранение дублирования, затем применить PreparedStatement, StringBuilder, вынос инвариантов, индексы и кэширование. На защите необходимо показать проект, классификацию ошибок, изменения, сохранение логики и измерение улучшений.")

add_heading(doc, "2 Анализ исходной ветки")
add_heading(doc, "2.1 Фактически найденные проблемы", 2)
actual_errors = [
    ("Функциональная", "Кнопка добавления студента вызывает handleAddSubject.", "Студент не добавляется."),
    ("Конфигурация", "В javafx-maven-plugin указан несуществующий главный класс.", "Команда javafx:run не запускает приложение."),
    ("Безопасность", "Учётные данные MySQL находятся в Java-коде и pom.xml.", "Секрет попадает в Git и расходится между конфигурациями."),
    ("Оптимизация", "Список студентов загружается дважды в loadData().", "Лишний запрос и повторное отображение тех же строк."),
    ("Оптимизация", "JDBC вызывается из JavaFX Application Thread.", "При медленной БД окно зависает."),
    ("Надёжность", "SQLException только печатается в консоль.", "Интерфейс не сообщает пользователю о сбое."),
    ("Валидация", "Оценка проверяется только как число.", "Можно сохранить значение вне диапазона 1–5."),
    ("Оптимизация", "Загружается полный список студентов.", "Память и время растут вместе с объёмом таблицы."),
    ("Ресурсы", "Статическое соединение закрывается каждым DAO-методом.", "Заявленное повторное использование соединения отсутствует."),
    ("База данных", "Нет отдельного индекса под агрегирование оценок.", "План запроса хуже масштабируется."),
    ("Рефакторинг", "Модели изменяемы и не проверяют состояние.", "Некорректный объект можно создать после конструктора."),
]
add_table(doc, ["Класс", "Наблюдение", "Последствие"], actual_errors, [3.2, 7.2, 6.1], 9)

add_heading(doc, "2.2 Сопоставление с перечнем группы 4", 2)
add_body(doc, "Полученная ветка уже содержит часть правильных решений. Поэтому наличие ошибок оценивалось по коду, а не только по названию ветки.")
checklist = [
    ("Рефакторинг", "N+1 и вложенный запрос", "Нет", "Использован один LEFT JOIN с AVG."),
    ("Рефакторинг", "Длинный getAllStudents", "Нет", "Вложенного цикла нет."),
    ("Рефакторинг", "if без else в showAverage", "Нет", "Метод showAverage отсутствует."),
    ("Рефакторинг", "Магические числа", "Частично", "100.0 используется для округления."),
    ("Рефакторинг", "Группы данных", "Частично", "Имя и группа передаются отдельно."),
    ("Рефакторинг", "Открытый grades", "Нет", "Коллекции grades в модели нет."),
    ("Рефакторинг", "Нет комментариев", "Нет", "Комментарии присутствуют, часть избыточна."),
    ("Оптимизация", "N+1", "Нет", "Среднее загружается одним запросом."),
    ("Оптимизация", "Нет PreparedStatement", "Нет", "Все пользовательские значения параметризованы."),
    ("Оптимизация", "Среднее O(n) в Java", "Нет", "Агрегирование выполняет SQL AVG."),
    ("Оптимизация", "Конкатенация INSERT", "Нет", "INSERT выполняются с параметрами."),
    ("Оптимизация", "Нет индекса student_id", "Частично", "Нет явного составного индекса под агрегирование."),
    ("Оптимизация", "Map оценок в памяти", "Нет", "Оценки хранятся в MySQL."),
    ("Оптимизация", "Нет буферизации ResultSet", "Частично", "Нет пагинации или fetch-size."),
]
add_table(doc, ["Блок", "Пункт методички", "В исходной ветке", "Основание"], checklist, [2.5, 5.0, 3.0, 6.0], 8)

add_heading(doc, "3 Выполненный рефакторинг")
changes = [
    "исправлена привязка кнопки добавления студента к handleAddStudent;",
    "точка входа Maven заменена на com.example.rpm_pr_1.Launcher;",
    "секреты удалены из исходников, настройки читаются из системных свойств или переменных окружения;",
    "классы Student и Subject сделаны неизменяемыми;",
    "связанные поля имени и группы выделены в StudentIdentity;",
    "диапазон оценки и округление выделены в GradeValue без магического числа 100;",
    "преобразование ResultSet вынесено в mapStudent и mapSubject;",
    "ошибки доступа к данным преобразуются в DataAccessException и показываются в интерфейсе.",
]
for item in changes:
    add_bullet(doc, item)

add_code(doc, [
    "public record GradeValue(double value) {",
    "    public static final double MIN_VALUE = 1.0;",
    "    public static final double MAX_VALUE = 5.0;",
    "    public GradeValue {",
    "        if (!Double.isFinite(value) || value < MIN_VALUE || value > MAX_VALUE)",
    "            throw new IllegalArgumentException(\"Оценка должна быть числом от 1 до 5\");",
    "    }",
    "}",
], "Листинг 1 – Объект значения и Introduce Constant")

add_heading(doc, "4 Выполненная оптимизация")
optimizations = [
    "DAO загружает страницу по 25 студентов через LIMIT и OFFSET;",
    "список студентов читается один раз и повторно используется таблицей и ComboBox;",
    "список предметов кэшируется между обычными обновлениями;",
    "операции БД выполняются последовательным фоновым ExecutorService через JavaFX Task;",
    "добавлены составные индексы grades(student_id, grade_value) и grades(subject_id, grade_value);",
    "сохранены PreparedStatement, LEFT JOIN, SQL AVG и try-with-resources, которые уже были правильными.",
]
for item in optimizations:
    add_bullet(doc, item)

add_code(doc, [
    "Task<ScreenData> task = new Task<>() {",
    "    @Override protected ScreenData call() {",
    "        int count = gradeDAO.countStudents();",
    "        List<Student> students = gradeDAO.getStudentsPage(PAGE_SIZE, page * PAGE_SIZE);",
    "        return new ScreenData(students, subjects, pages, page);",
    "    }",
    "};",
    "databaseExecutor.execute(task);",
], "Листинг 2 – Выполнение JDBC вне JavaFX Application Thread")

add_heading(doc, "5 Доказательство сохранения логики")
add_heading(doc, "5.1 Автоматизированная проверка", 2)
add_body(doc, "Предметные операции не удалялись: приложение по-прежнему добавляет студентов и предметы, сохраняет оценки и показывает средний балл. Запрос среднего значения сохранил LEFT JOIN и AVG, поэтому студент без оценок остаётся в таблице, а среднее вычисляется базой данных.")
add_body(doc, "Модельный smoke-тест создал студента, проверил неизменность имени и среднего балла и подтвердил отклонение оценки 6 и пустого имени. Результат: MODEL_SMOKE_TEST_OK. Проверка FXML подтвердила наличие всех пяти обработчиков, XML корректен. Компиляция десяти основных исходных файлов завершилась с кодом 0 и создала 13 class-файлов.")

add_heading(doc, "5.2 Ручная проверка с MySQL", 2)
add_body(doc, "После создания схемы student_db и настройки параметров подключения приложение успешно запущено из IntelliJ IDEA. Проверены загрузка существующих студентов, добавление новой записи, загрузка предметов и сохранение оценки. В таблице отображаются рассчитанные средние баллы, а строка состояния после фонового запроса показывает «Готово».")
add_body(doc, "Исправленная версия отображает по 25 записей на странице и корректно блокирует кнопки перехода, когда доступна только одна страница. Пустой результат запроса также обрабатывается без исключения: таблица остаётся доступной, элементы управления не перекрываются, индикатор состояния сообщает о завершении операции.")

for index, (image_path, alt_text, caption_text) in enumerate(SCREENSHOTS):
    if image_path.exists():
        if index > 0:
            doc.add_page_break()
        add_figure(doc, image_path, alt_text, caption_text)

add_heading(doc, "6 Измерение улучшений")
add_body(doc, "Функциональный сценарий проверен на локальной MySQL. Для оценки масштабирования дополнительно рассчитан объём обработки при 1000 студентах: исходный loadData() отображает список дважды и выполняет до 2000 преобразований строк Student. Исправленная версия получает только текущую страницу из 25 строк, поэтому объём отображения уменьшается в 80 раз.")
metrics = [
    ("Обычное обновление после загрузки предметов", "3 запроса", "2 запроса", "на 33% меньше обращений"),
    ("Строки Student при 1000 записях", "до 2000", "не более 25", "в 80 раз меньше"),
    ("JDBC-операции в JavaFX-потоке", "все", "0", "интерфейс не блокируется"),
    ("Учётные данные в отслеживаемых файлах", "2 места", "0", "секреты вынесены из Git"),
]
add_table(doc, ["Показатель", "До", "После", "Результат"], metrics, [6.1, 2.7, 2.7, 5.0], 9)

add_heading(doc, "7 Оценка исходного проекта")
add_body(doc, "Методичка не содержит готовой балльной шкалы. Для единообразной оценки введены четыре критерия общей стоимостью 100 баллов.")
scores = [
    ("Функциональная корректность", "15/30", "Интерфейс существует, но добавление студента и Maven-запуск настроены неверно."),
    ("Качество кода", "17/25", "Слои выделены, но отсутствуют доменная валидация и корректная передача ошибок."),
    ("БД и производительность", "15/25", "JOIN и PreparedStatement есть, но остаются повторный запрос, UI-поток и отсутствие пагинации."),
    ("Git и доказательства", "5/20", "Ветка с ошибками есть, но master не содержит полного рефакторинга, отчёта и измерений."),
    ("Итого", "52/100", "Проект частично готов, этапы 3–5 методички выполнены не полностью."),
]
add_table(doc, ["Критерий", "Балл", "Обоснование"], scores, [5.0, 2.5, 9.0], 10)

add_heading(doc, "7.1 Оценка исправленной версии", 2)
fixed_scores = [
    ("Функциональная корректность", "30/30", "Добавление студента, предмета и оценки проверено на MySQL; средний балл отображается."),
    ("Качество кода", "23/25", "Добавлены доменная валидация, неизменяемые модели, понятная передача ошибок и отдельные объекты значений."),
    ("БД и производительность", "23/25", "Запросы вынесены из UI-потока, добавлены пагинация, кэширование и индексы; крупный нагрузочный тест не проводился."),
    ("Git и доказательства", "18/20", "Исходная и исправленная ветки разделены, сборка и ручной сценарий подтверждены, приложены скриншоты."),
    ("Итого", "94/100", "Требования анализа, рефакторинга, оптимизации и проверки выполнены."),
]
add_table(doc, ["Критерий", "Балл", "Обоснование"], fixed_scores, [5.0, 2.5, 9.0], 10)

add_heading(doc, "8 Материалы для аргументации")
sources = [
    "Martin Fowler. Refactoring: https://refactoring.com",
    "Oracle. Using Prepared Statements: https://docs.oracle.com/javase/tutorial/jdbc/basics/prepared.html",
    "OpenJFX. Task: https://openjfx.io/javadoc/17/javafx.graphics/javafx/concurrent/Task.html",
    "MySQL. How MySQL Uses Indexes: https://dev.mysql.com/doc/refman/8.0/en/mysql-indexes.html",
    "OWASP. Secrets Management: https://cheatsheetseries.owasp.org/cheatsheets/Secrets_Management_Cheat_Sheet.html",
]
for item in sources:
    add_bullet(doc, item)

add_heading(doc, "Заключение")
add_body(doc, "В ходе работы проанализирован чужой JavaFX-проект из ветки «ошибки-человечества». Название ветки само по себе не использовалось как доказательство: проверка показала, что несколько ошибок из перечня группы 4 уже отсутствуют. Исходная версия применяет LEFT JOIN, AVG, PreparedStatement и try-with-resources, поэтому эти решения сохранены и отмечены как корректные.")
add_body(doc, "Реально обнаружены ошибки привязки FXML, точки входа Maven, хранения учётных данных, повторной загрузки студентов, обращения к БД из UI-потока, обработки исключений, валидации и масштабирования. В отдельной ветке они исправлены с помощью инкапсуляции, Extract Class, Introduce Constant, выделения методов отображения, пагинации, кэширования, индексов и JavaFX Task.")
add_body(doc, "Предметная логика не нарушена, что подтверждается успешной Maven-сборкой, модельным smoke-тестом, проверкой FXML и ручным запуском с MySQL. В интерфейсе проверены чтение списка студентов, добавление данных, сохранение оценки, расчёт среднего балла, пагинация и пустой результат. Исходный проект оценён в 52 балла из 100, исправленная версия — в 94 балла из 100. Разница объясняется устранением функциональных ошибок, безопасной настройкой подключения, переносом JDBC-операций из JavaFX-потока и добавлением доказательств проверки.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(footer.add_run("Отчёт по анализу проекта RPM_PR_1"), size=9, color="666666")

doc.save(OUTPUT)
print(OUTPUT)
